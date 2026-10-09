"""Built to Share guest pipeline: every B2S list into People, Organisations and Engagements, as one Notion AI prompt.

Usage:
  python migration/b2s_pipeline.py          parse, dedupe, match against Notion, write the prompt, review CSV and counts

Reads (git-ignored, data/raw/b2s/):
  b2s_guest_list_2026-10-09.xlsx            export of Google Sheet "B2S Guest list" (Chaandini), 10 tabs
  notion_live_2026-10-09.json               the live Notion "Built-2-Share" table (133 rows)
  notion_archive_2026-10-09.json            the trashed earlier copy of it (492 rows)
  notion_people_2026-10-09.json             People: id|name|tags|has_email
  notion_people_contacts_2026-10-09.json    People with an email or LinkedIn: id|email|linkedin|other emails|org id
  notion_orgs_2026-10-09.json               Organisations: id|name|domain|website
  notion_b2s_engagements_2026-10-09.json    the 19 Built to Share Engagements already in Notion
Writes:
  data/out/b2s_pipeline/candidates.json     one record per person, every source row it came from, mapped fields
  data/out/b2s_pipeline/notion_ai_b2s_pipeline.md   the one prompt Akshay uploads to Notion AI (decision 91)
  data/review/b2s_pipeline_review.csv       rows settled by a rule worth a second look, and the few left open

How B2S works and the mapping were agreed with Akshay on 9 Oct (decision 95). Satish's approval is Host verdict;
the bucket he approves by is an Engagement tag; Tier is pre-filled; the 2-month retry is the Next date on
Unresponsive rows. Existing Engagements are never regressed and only their empty fields are filled.
Never edits the raw files. Re-run safely.
"""
import collections
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paas import EMAIL, LEGAL, fold, li_slug, li_url, name_key  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data/raw/b2s"
XLSX = RAW / "b2s_guest_list_2026-10-09.xlsx"
OUT = ROOT / "data/out/b2s_pipeline"
REVIEW = ROOT / "data/review/b2s_pipeline_review.csv"
TODAY = dt.date(2026, 10, 9)
BATCH = "b2s-pipeline-2026-10-09"
SHEET = "B2S Guest list"

PEOPLE_DB = "https://app.notion.com/p/f54094b86c3549dead74c0034afad86c"
ORGS_DB = "https://app.notion.com/p/4e484b1c510b42ce9ea979307bff31e4"
ENGS_DB = "https://app.notion.com/p/5e77c373909a4bf6b9195570de46c63b"
SHOW = "3ed19a301a7f819e8f45de04a7c9d025"          # Built to Share
SATISH = "3ed19a301a7f81c59b10d48cb642a3bf"        # Satish Mugulavalli, Hissa, the B2S host

ACADEMIC = re.compile(r"\.(ac\.[a-z]{2}|edu(\.[a-z]{2})?)$")
FREE_MAIL = {"gmail.com", "yahoo.com", "outlook.com", "hotmail.com", "icloud.com", "me.com", "proton.me", "protonmail.com"}
GREEN = {"FF00FF00", "FFB6D7A8", "FF93C47D"}       # Satish approved (India tab, N2)
YELLOW = {"FFFFFF00"}                               # Global known in India: Satish reaches them himself
AD_FILL = {"FFFFF2CC"}                              # N2: names Akshay knows (Mu Sigma, Zeta, Moglix)

# Status ladder rank: the furthest-along status across sources wins
RANK = {"New": 0, "Contacted": 1, "Unresponsive": 2, "LT chase": 3, "They declined": 4, "We declined": 4,
        "Booking": 5, "Booked": 6, "Recorded": 7, "Form filled": 8}
OPEN = {"New", "Contacted", "LT chase", "Booking", "Booked"}
NOTION_STATUS = {"not started": "New", "emailed / started": "Contacted", "chasing": "Contacted",
                 "maybe later": "LT chase", "approved": "Booking", "scheduled": "Booked", "booked": "Booked",
                 "recorded": "Recorded", "released": "Form filled", "unresponsive": "Unresponsive",
                 "declined": "They declined", "do not contact": "We declined"}
SHEET_STATUS = {"declined": "They declined", "unresponsive": "Unresponsive", "chasing": "Contacted",
                "maybe later": "LT chase", "started": "Contacted", "approved": "Booking"}
BUCKET = {"outliers": "Outlier", "outlier": "Outlier", "recently ipod": "Recently IPO'd",
          "recently ipo d": "Recently IPO'd", "ipo pre ipo stage": "Pre-IPO", "pre ipo ipo": "Pre-IPO",
          "big indian founders": "Big Indian founder", "big global founders": "Big global founder",
          "more of the same": "More of the same", "foreign unicorns": "Global unicorn", "global unicorn": "Global unicorn"}
TIER_15 = {"Outlier", "Recently IPO'd", "Pre-IPO", "Big Indian founder", "Big global founder"}
N2_INVALID = {"DEPARTED": "founder has left the company", "COMPANY_DEAD": "company shut down", "BAD_ROW": "no usable contact"}
# Dates said in words, from Akshay's comments in the live table
LT_DATES = {"jaydeep barman": dt.date(2027, 1, 5), "sanjeev bikhchandani": dt.date(2026, 11, 2)}
# Rows the parser cannot settle alone, fixed by research (source id -> value)
NAME_FIX = {"rohit": "Rohit Chennamaneni"}                 # DarwinBox: full name from his email
LIVE_TO_ENG = {                                            # Released rows in the live table -> Engagement from Blue
    "fibe": "3ed19a301a7f81448df2f0c4ce15110f", "brightchamps": "3ed19a301a7f81fea7b7c41f40fda566",
    "mudrex": "3ed19a301a7f81a4b953ccdc8ff4347c", "idfy": "3ed19a301a7f81cdb594dbe9917ae47a",
    "wonderchef": "3ed19a301a7f8140a4d9f16653cf62bb", "cashify": "3ed19a301a7f81e68034fa7bf4c8bcd0",
    "everpure pure storage": "3ed19a301a7f8116954aff0a31cab224", "arya ai": "3ed19a301a7f817dbe8adb7443eca181",
    "1buy ai": "3ed19a301a7f8153b8a0f79f10c99a8f", "share market": "3ed19a301a7f8137b61ef2887be8accc",
    "cardekho": "3ed19a301a7f81809e4ad9a8d22275fb", "ideaforge": "3ed19a301a7f81bcadefce7d76bdcdfe",
    "asia healthcare holdings": "3ed19a301a7f8194bc07cea45403b4ec", "betterplace": "3ed19a301a7f81cba8f6dbe4964cc6ab",
    "scripbox": "3ed19a301a7f81b991b0d506fab930e4", "jaipur rugs": "3ed19a301a7f81b98190fa1d90253955",
    "cheq": "3f119a301a7f8196ac99de24b5dd4797", "appsforbharat": "3f119a301a7f819f9a19f9f8a97ddb3e",
}
# People rows matched by research where the name alone is ambiguous or first-name only
PERSON_FIX = {
    "anurag jain|cardekho": "3ed19a301a7f81f9a00bea5f67b9d562",   # the Guest row, org CarDekho (the other is a US VC)
    "ankit agrawal|insurancedekho": "3ed19a301a7f81f9974af5be0a9f0b2f",  # the row linked to InsuranceDekho
}
# Same name as a People row but a different person (checked): a new row, flagged Possible duplicate for the dedupe check
DIFFERENT_PERSON = {"jon miller|halcyon", "amit jain|luma ai", "alex kinnier|upside", "andrew parker|papa"}


def val(c):
    return re.sub(r"\s+", " ", str(c.value)).strip() if c is not None and c.value is not None else ""


def fill_of(c):
    f = c.fill
    if f is not None and f.fgColor is not None and f.fgColor.type == "rgb" and f.fgColor.rgb not in ("00000000", "FFFFFFFF"):
        return f.fgColor.rgb
    return None


def date_of(c):
    v = c.value if c is not None else None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def split_name(cell_text):
    """'Name @ Company' or 'Name@Company' -> (name, company)."""
    t = cell_text.strip()
    if "@" in t and not EMAIL.fullmatch(t):
        n, c = t.split("@", 1)
        return n.strip(), c.strip() or None
    return t, None


def split_people(name):
    """'Varun & Ghazal Alagh' -> ['Varun Alagh', 'Ghazal Alagh']; 'Roy Mann & Eran Zinman' -> both."""
    name = re.sub(r"\s*\(.*?\)\s*", " ", name).strip()
    parts = [p.strip() for p in re.split(r"\s+&\s+|\s+and\s+", name) if p.strip()]
    if len(parts) == 2 and len(parts[0].split()) == 1 and len(parts[1].split()) > 1:
        parts[0] = f"{parts[0]} {parts[1].split()[-1]}"
    return parts


def titlecase(n):
    return " ".join(w.capitalize() if w.isupper() or w.islower() else w for w in n.split())


def org_name_key(s):
    s = fold(re.sub(r"\(.*?\)", " ", str(s or "")))
    return re.sub(r"\s+", " ", LEGAL.sub(" ", s)).strip() or s


def email_info(text):
    """Founder email from a cell. A dagger marks an address that belongs to someone else (Global p3)."""
    t = str(text or "")
    found = [e.lower().strip(".") for e in EMAIL.findall(t)]
    if not found:
        return None, None
    if "†" in t:
        return None, found[0]
    return found[0], None


def rec(src, sid, name, company):
    return {"src": src, "sid": sid, "name": name, "company": company, "email": None, "other_email": None, "li": None,
            "title": None, "country": None, "valuation": None, "bucket": None, "connect": None, "status": None,
            "raw_status": None, "sent": [], "next": None, "approved": None, "satish_reach": False, "notes": [],
            "flag": None}


def follow_dates(cells):
    sent, notes = [], []
    for c in cells:
        d = date_of(c)
        if d:
            sent.append(d)
        elif val(c):
            notes.append(val(c))
    return sent, notes


def read_sheet():
    wb = openpyxl.load_workbook(XLSX)
    out, skipped = [], []

    ws = wb["AD connect"]
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        if not val(r[0]):
            continue
        x = rec("AD connect", f"{SHEET}/AD connect/{i}", val(r[0]), val(r[1]))
        x.update(connect="AD", bucket="Big Indian founder" if val(r[2]) == "Indian" else "Big global founder")
        x["notes"].append(f"FT link: {val(r[3])}")
        out.append(x)

    ws = wb["Approved By Satish"]
    section = None
    for i, r in enumerate(ws.iter_rows(min_row=3), start=3):
        if not val(r[0]):
            continue
        if not val(r[1]):
            section = BUCKET.get(fold(val(r[0])))
            continue
        x = rec("Approved By Satish", f"{SHEET}/Approved By Satish/{i}", val(r[0]), val(r[1]))
        x["email"], x["other_email"] = email_info(val(r[2]))
        x["sent"], notes = follow_dates(r[3:8])
        for n in notes:
            if n.lower() == "ad connect":
                x["connect"] = "AD"
            elif n.lower() == "satish connect":
                x["connect"] = "Satish"
            else:
                x["notes"].append(n)
        x.update(approved="Approved By Satish tab", bucket=section)
        out.append(x)

    ws = wb["Global Unicorn"]
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        if not val(r[0]):
            continue
        n, c = split_name(val(r[0]))
        x = rec("Global Unicorn", f"{SHEET}/Global Unicorn/{i}", n, c)
        x["email"], x["other_email"] = email_info(val(r[2]))
        x.update(title=val(r[1]) or None, country=val(r[3]) or None, valuation=val(r[4]) or None,
                 bucket="Global unicorn", raw_status=val(r[5]) or None)
        out.append(x)

    ws = wb["India"]
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        if not val(r[0]):
            continue
        n, c = split_name(val(r[0]))
        x = rec("India", f"{SHEET}/India/{i}", n, c)
        conn = val(r[1]).lower()
        channel = val(r[5])
        x["email"], _ = email_info(channel)
        x["sent"], notes = follow_dates(r[7:])
        fill = fill_of(r[0])
        x.update(connect="AD" if conn == "akshay" else None, raw_status=val(r[2]) or None, valuation=val(r[3]) or None,
                 bucket="Big Indian founder", approved="green on the India tab" if fill in GREEN else None)
        if channel and not x["email"]:
            x["notes"].append(f"Outreach channel: {channel}")
        if "hissa" in channel.lower():
            x["connect"] = "Satish"
        x["notes"] += [f"Note: {n}" for n in notes]
        out.append(x)

    ws = wb["Global known in India"]
    for i, r in enumerate(ws.iter_rows(min_row=3), start=3):
        if not val(r[0]):
            continue
        n, c = split_name(val(r[0]))
        x = rec("Global known in India", f"{SHEET}/Global known in India/{i}", n, c)
        x["email"], x["other_email"] = email_info(val(r[7]))
        x["sent"], notes = follow_dates(r[9:])
        fill = fill_of(r[0])
        x.update(raw_status=val(r[2]) or None, country=val(r[3]) or None, valuation=val(r[4]) or None,
                 bucket="Big global founder", approved="Global known in India tab (default agreed 9 Oct)",
                 satish_reach=fill in YELLOW)
        if fill in YELLOW:
            x["connect"] = "Satish"
        x["notes"] += [f"Note: {n}" for n in notes]
        out.append(x)

    ws = wb["NEW"]
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        raw = str(r[0].value or "").strip()
        if not raw:
            continue
        m = re.match(r'=HYPERLINK\("([^"]+)"\s*,\s*"([^"]+)"\)', raw)
        li, text = (m.group(1), m.group(2)) if m else (None, raw)
        if "@" not in text:
            skipped.append([f"{SHEET}/NEW/{i}", text, "section label or name with no company; the person is on another tab"])
            continue
        n, c = split_name(text)
        x = rec("NEW", f"{SHEET}/NEW/{i}", n, c)
        x["email"], x["other_email"] = email_info(val(r[2]))
        x["sent"], _ = follow_dates(r[5:10])
        x.update(li=li, country=val(r[1]) or None, raw_status=val(r[3]) or None, bucket="Global unicorn")
        out.append(x)

    ws = wb["Outliers"]
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        if not val(r[0]):
            continue
        x = rec("Outliers", f"{SHEET}/Outliers/{i}", val(r[0]), val(r[1]))
        x["email"], x["other_email"] = email_info(val(r[2]))
        x["sent"], _ = follow_dates(r[3:8])
        x["bucket"] = "Outlier"
        out.append(x)

    ws = wb["N2"]
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        if not val(r[0]):
            continue
        x = rec("N2", f"{SHEET}/N2/{i}", val(r[0]), val(r[1]))
        x["email"], x["other_email"] = email_info(val(r[2]))
        x["sent"], _ = follow_dates(r[6:11])
        fill = fill_of(r[0])
        x.update(flag=val(r[3]) or None, raw_status=val(r[4]) or None, bucket="Global unicorn",
                 approved="green on the N2 tab" if fill in GREEN else None, connect="AD" if fill in AD_FILL else None)
        out.append(x)

    ws = wb["N3 (6)"]
    for i, r in enumerate(ws.iter_rows(min_row=4), start=4):
        if not val(r[0]):
            continue
        n, c = split_name(val(r[0]))
        x = rec("N3", f"{SHEET}/N3 (6)/{i}", n, c)
        x["email"], x["other_email"] = email_info(val(r[2]))
        x["sent"], _ = follow_dates(r[5:10])
        x.update(bucket=BUCKET.get(fold(val(r[1]))), raw_status=val(r[3]) or None)
        if val(r[4]):
            x["notes"].append(val(r[4]))
        out.append(x)

    ws = wb["Global p3"]
    for i, r in enumerate(ws.iter_rows(min_row=3), start=3):
        if not val(r[0]):
            continue
        x = rec("Global p3", f"{SHEET}/Global p3/{i}", val(r[0]), val(r[1]))
        x["email"], x["other_email"] = email_info(val(r[6]))
        x.update(raw_status=val(r[3]) or None, country=val(r[4]) or None, valuation=val(r[5]) or None,
                 bucket="Global unicorn", connect="AD" if val(r[2]).lower() in ("ad", "akshay") else None)
        out.append(x)
    return out, skipped


def read_notion(path, label):
    out, skipped = [], []
    for r in json.load(open(path)):
        name = r["name"] or ""
        if name.upper().startswith("DELETE"):
            skipped.append([f"{label}/{r['id']}", name, "marked DELETE ME in the source table: a known duplicate"])
            continue
        n, c = split_name(name)
        x = rec(label, f"{label}/{r['id']}", n, c)
        x["email"], _ = email_info(r.get("email"))
        cmt = r.get("cmt") or ""
        x.update(bucket=BUCKET.get(fold(r.get("cat") or "")), country=r.get("country"),
                 raw_status=r.get("status"), connect={"AD": "AD", "Satish": "Satish"}.get(r.get("connect")),
                 approved="Satish Approved in the Notion table" if r.get("appr") == "Yes" or "satish approved" in cmt.lower() else None)
        if r.get("next"):
            x["next"] = dt.date.fromisoformat(r["next"][:10])
        if cmt:
            x["notes"].append(f"AD comment: {cmt}")
        out.append(x)
    return out, skipped


def status_of(x):
    raw = (x["raw_status"] or "").strip().lower()
    table = NOTION_STATUS if x["src"].startswith("Notion") else SHEET_STATUS
    s = table.get(raw)
    if s is None and raw:
        return None
    if s is None:
        s = "Contacted" if x["sent"] else "New"
    return s


def main():
    sheet_rows, skipped = read_sheet()
    live, sk1 = read_notion(RAW / "notion_live_2026-10-09.json", "Notion B2S")
    arch, sk2 = read_notion(RAW / "notion_archive_2026-10-09.json", "Notion B2S archive")
    skipped += sk1 + sk2
    rows = live + sheet_rows + arch           # live table first: it is the most current record
    review = list(skipped)

    for x in rows:
        x["status"] = status_of(x)
        if x["status"] is None:
            review.append([x["sid"], x["name"], f"unmapped status {x['raw_status']!r}, read as New", ""])
            x["status"] = "New"
        x["people"] = [titlecase(p) for p in split_people(x["name"])]
        if len(x["people"]) == 1 and len(x["people"][0].split()) == 1:
            fixed = NAME_FIX.get(fold(x["people"][0]))
            if fixed:
                x["people"] = [fixed]

    # Group source rows into people: email, then LinkedIn, then name
    groups, by = [], {}
    for x in rows:
        if re.match(r"no founder", x["name"], re.I):
            x["people"] = [""]
        if not x["people"] or not x["people"][0]:
            review.append([x["sid"], x["name"], "no founder named: Organisation and an Engagement with no Guest, "
                           "Next action asks Satish for the name", x["company"] or ""])
            groups.append({"rows": [x], "nobody": True})
            continue
        keys = []
        if x["email"]:
            keys.append("em:" + x["email"])
        if li_slug(x["li"]):
            keys.append("li:" + li_slug(x["li"]))
        keys.append("nm:" + name_key(x["people"][0]))
        g = next((by[k] for k in keys if k in by), None)
        if g is None:
            g = {"rows": []}
            groups.append(g)
        g["rows"].append(x)
        for k in keys:
            by.setdefault(k, g)

    people = json.load(open(RAW / "notion_people_2026-10-09.json"))
    P_name = collections.defaultdict(list)
    P_tags = {}
    for line in people:
        pid, nm, tags, _ = (line.split("|") + ["", "", ""])[:4]
        P_name[name_key(nm)].append(pid)
        P_tags[pid] = tags
    P_email, P_li, P_org, P_has = {}, {}, {}, {}
    for line in json.load(open(RAW / "notion_people_contacts_2026-10-09.json")):
        pid, em, li, other, org = (line.split("|") + ["", "", "", ""])[:5]
        for e in [em] + EMAIL.findall(other):
            if e:
                P_email.setdefault(e.lower(), pid)
        if li_slug(li):
            P_li.setdefault(li_slug(li), pid)
        P_org[pid] = org
        P_has[pid] = (bool(em), bool(li))
    orgs = json.load(open(RAW / "notion_orgs_2026-10-09.json"))
    O_name, O_dom = collections.defaultdict(list), {}
    O_label = {}
    for line in orgs:
        oid, nm, dom, web = (line.split("|") + ["", "", ""])[:4]
        O_name[org_name_key(nm)].append(oid)
        O_label[oid] = nm
        for d in (dom, re.sub(r"^https?://(www\.)?", "", web).split("/")[0]):
            if d:
                O_dom.setdefault(d.lower(), oid)
    existing = {e["id"]: e for e in json.load(open(RAW / "notion_b2s_engagements_2026-10-09.json"))}
    eng_by_guest = {e["guest"]: e["id"] for e in existing.values() if e["guest"]}

    new_orgs, cands = {}, []
    for g in groups:
        xs = g["rows"]
        head = xs[0]
        c = {"people": max((x["people"] for x in xs), key=lambda p: (len(p), len(" ".join(p)))),
             "source_ids": [x["sid"] for x in xs], "tabs": list(dict.fromkeys(x["src"] for x in xs))}
        c["name"] = c["people"][0]
        c["nobody"] = bool(g.get("nobody"))
        for f in ("company", "email", "other_email", "li", "title", "country", "valuation", "connect", "approved"):
            c[f] = next((x[f] for x in xs if x[f]), None)
        # Bucket: Satish's own section first, then N3's bucket call, then the Notion category, then the tab
        order = ["Approved By Satish", "N3", "Notion B2S", "Notion B2S archive", "Outliers", "AD connect", "India",
                 "Global known in India", "NEW", "N2", "Global Unicorn", "Global p3"]
        c["bucket"] = next((x["bucket"] for o in order for x in xs if x["src"] == o and x["bucket"]), None)
        c["satish_reach"] = any(x["satish_reach"] for x in xs)
        c["sent"] = sorted({d for x in xs for d in x["sent"]})
        c["next_src"] = max((x["next"] for x in xs if x["next"]), default=None)
        c["notes"] = list(dict.fromkeys(n for x in xs for n in x["notes"]))
        c["flags"] = sorted({x["flag"] for x in xs if x["flag"] and x["flag"] != "CURRENT"})
        statuses = [x["status"] for x in xs]
        best = max(statuses, key=lambda s: RANK[s])
        recent = [d for d in c["sent"] if d >= dt.date(2026, 10, 1)]
        if best == "Unresponsive" and recent:
            best = "Contacted"                 # a second campaign has started since they went quiet
        done = [d for d in c["sent"] if d <= TODAY]
        if best == "Contacted" and len(done) >= 5 and not [d for d in c["sent"] if d > TODAY] \
                and done[-1] <= TODAY - dt.timedelta(days=7):
            best = "Unresponsive"              # the 5-mail sequence ran out with no reply
            c["notes"].insert(0, "Sequence finished with no reply.")
        c["status"] = best
        c["raw_statuses"] = sorted({f"{x['src']}: {x['raw_status']}" for x in xs if x["raw_status"]})
        if set(c["flags"]) & {"DEPARTED", "COMPANY_DEAD", "DOMAIN_CHANGED", "BAD_ROW"}:
            c["stale_email"], c["email"] = c["email"], None   # the N2 check says this address no longer reaches them
        invalid = [N2_INVALID[f] for f in c["flags"] if f in N2_INVALID]
        if invalid and c["status"] == "New":
            c["status"] = "We declined"
            c["declined_reason"] = invalid[0]
        cands.append(c)

    # Match People
    for c in cands:
        pid, how = None, None
        key = f"{name_key(c['name'])}|{fold(c['company'] or '')}"
        if key in PERSON_FIX:
            pid, how = PERSON_FIX[key], "research"
        if not pid and c["email"] and c["email"] in P_email:
            pid, how = P_email[c["email"]], "email"
        if not pid and li_slug(c["li"]) and li_slug(c["li"]) in P_li:
            pid, how = P_li[li_slug(c["li"])], "LinkedIn"
        c["dup_of"] = None
        if not pid and not c["nobody"] and len(c["name"].split()) > 1:
            hits = P_name.get(name_key(c["name"]), [])
            if key in DIFFERENT_PERSON or len(hits) > 1:
                c["dup_of"] = hits[0] if hits else None
                review.append([c["source_ids"][0], c["name"], "same name as an existing People row but not the same person: new row, "
                               "flagged Possible duplicate", ", ".join(hits)])
            elif len(hits) == 1:
                pid, how = hits[0], "name"
        c["person_id"], c["match"] = pid, how
        c["person_tags"] = P_tags.get(pid, "") if pid else ""
        c["fill_email"] = bool(pid and c["email"] and not P_has.get(pid, (False, False))[0] and c["email"] not in P_email)
        c["fill_li"] = bool(pid and li_url(c["li"]) and not P_has.get(pid, (False, False))[1])

    # Match or create Organisations
    for c in cands:
        already = c["person_id"] in eng_by_guest or (fold(c["company"] or "") in LIVE_TO_ENG and c["status"] in ("Booked", "Recorded", "Form filled"))
        if not c["company"] or already:
            c["org_id"] = c["org_new"] = None   # an existing Engagement already carries its Organisation
            continue
        ok = org_name_key(c["company"])
        dom = c["email"].split("@")[1] if c["email"] else None
        if dom and ACADEMIC.search(dom):
            dom = None
        oid = None
        if c["person_id"] and P_org.get(c["person_id"]):
            oid = P_org[c["person_id"]] if org_name_key(O_label.get(P_org[c["person_id"]], "")) == ok else None
        if not oid and len(O_name.get(ok, [])) >= 1:
            oid = O_name[ok][0]
        if not oid and dom and dom not in FREE_MAIL and dom in O_dom:
            oid = O_dom[dom]
        c["org_id"] = oid
        c["org_new"] = None
        if not oid:
            k = ok
            if k not in new_orgs:
                new_orgs[k] = {"key": f"B2S org {len(new_orgs) + 1:03d}", "name": re.sub(r"\s*\(.*?\)\s*", " ", c["company"]).strip(),
                               "domain": dom if dom and dom not in FREE_MAIL else None, "source_id": c["source_ids"][0]}
            elif not new_orgs[k]["domain"] and dom and dom not in FREE_MAIL:
                new_orgs[k]["domain"] = dom
            c["org_new"] = new_orgs[k]["key"]

    # Engagement fields
    for c in cands:
        c["eng_id"] = None
        if c["person_id"] and c["person_id"] in eng_by_guest:
            c["eng_id"] = eng_by_guest[c["person_id"]]
        live_key = fold(c["company"] or "")
        if not c["eng_id"] and live_key in LIVE_TO_ENG and c["status"] in ("Booking", "Booked", "Recorded", "Form filled"):
            c["eng_id"] = LIVE_TO_ENG[live_key]
        host = "Yes" if c["approved"] else "Not asked"
        tier = "Tier 1" if (c["approved"] or c["connect"]) else ("Tier 1.5" if c["bucket"] in TIER_15 else "Tier 2")
        tags = [c["bucket"]] if c["bucket"] else []
        if c["connect"]:
            tags.append("Warm intro")
        if not c["email"]:
            tags.append("No email")
        s = c["status"]
        last = c["sent"][-1] if c["sent"] else None
        nxt, action = None, None
        if s == "Contacted":
            future = [d for d in c["sent"] if d > TODAY]
            nxt = min(future) if future else max(c["next_src"] or TODAY, (last or TODAY) + dt.timedelta(days=7))
            action = f"Follow-up {min(len([d for d in c['sent'] if d <= TODAY]) + 1, 6)}"
        elif s == "LT chase":
            nxt = LT_DATES.get(name_key(c["name"])) or max(TODAY + dt.timedelta(days=30), (last or TODAY) + dt.timedelta(days=60))
            action = "Try again: they said later"
        elif s == "Unresponsive":
            nxt = (last or TODAY) + dt.timedelta(days=60)
            action = "Re-engage: 2 months since the last mail"
        elif s == "New":
            action = "Ask Satish" if host == "Not asked" else "First mail"
        waiting = []
        if s in OPEN and s not in ("Booking", "Booked"):
            if s == "New" and c["satish_reach"]:
                waiting = ["Satish"]
            elif c["connect"] == "AD" and s in ("New", "Contacted"):
                waiting = ["Akshay"]
            elif c["connect"] == "Satish" and s == "New":
                waiting = ["Satish"]
            else:
                waiting = ["Chaandini"]
        if c["satish_reach"] and s in ("New", "Unresponsive"):
            action, waiting = "Satish to reach out himself", ["Satish"]
        if c["nobody"]:
            action, waiting = "Satish to name the founder he knows here", ["Satish"]
        short = {"Notion B2S": "Notion", "Notion B2S archive": "Notion archive"}
        note = [f"B2S migration 9 Oct, from {', '.join(short.get(t, t) for t in c['tabs'])}."]
        was = sorted({x.split(": ", 1)[1] for x in c["raw_statuses"]})
        if was:
            note.append(f"Was: {', '.join(was)}.")
        if c["sent"]:
            note.append(f"Mails sent: {len([d for d in c['sent'] if d <= TODAY])}, last {last:%d %b}.")
        if c["approved"]:
            note.append(f"Satish approval: {c['approved']}.")
        if c["flags"]:
            note.append(f"N2 check: {', '.join(f.lower().replace('_', ' ') for f in c['flags'])}.")
        if c.get("declined_reason"):
            note.append(f"Not pursued: {c['declined_reason']}.")
        if c["other_email"]:
            note.append("Sheet email is someone else's, not used.")
        if c["valuation"]:
            note.append(f"Valuation {c['valuation']}.")
        note += [n[:160] for n in c["notes"] if "@" not in n][:2]
        c.update(host=host, tier=tier, eng_tags=tags, next_date=nxt.isoformat() if nxt else None, next_action=action,
                 waiting=waiting, chase=len([d for d in c["sent"] if d <= TODAY]) or None, note=" ".join(note)[:500])

    # Second founders on one row become their own People rows, Stakeholders on the Engagement
    extra = []
    cand_by_name = {name_key(c["name"]): c for c in cands if not c["nobody"]}
    for c in cands:
        c["stakeholders"] = []
        for p in c["people"][1:]:
            hits = P_name.get(name_key(p), [])
            other = cand_by_name.get(name_key(p))
            if other is not None:
                c["stakeholders"].append(("id", other["person_id"]) if other["person_id"] else ("new", other["source_ids"][0]))
            elif len(hits) == 1:
                c["stakeholders"].append(("id", hits[0]))
            else:
                extra.append({"name": p, "org_new": c["org_new"], "org_id": c["org_id"], "source_id": c["source_ids"][0] + " (co-founder)"})
                c["stakeholders"].append(("new", c["source_ids"][0] + " (co-founder)"))

    for i, c in enumerate(cands, 1):
        c["key"] = c["source_ids"][0]
        if c["eng_id"] is None and existing and c["person_id"] and c["person_id"] in eng_by_guest:
            c["eng_id"] = eng_by_guest[c["person_id"]]
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "candidates.json").write_text(json.dumps({"cands": cands, "orgs": new_orgs, "extra": extra}, indent=1, default=str))
    write_prompt(cands, new_orgs, extra, existing)

    for c in cands:
        if c["match"] == "name" and c["company"] and P_org.get(c["person_id"]) and c["org_id"] != P_org[c["person_id"]]:
            review.append([c["key"], c["name"], "matched People by name; Organisation named differently there, checked: same person",
                           f"B2S: {c['company']}; People org: {O_label.get(P_org[c['person_id']], '?')}"])
        if c["status"] == "We declined" and c.get("declined_reason"):
            review.append([c["key"], c["name"], f"We declined: {c['declined_reason']} (N2 check)", c["company"] or ""])
    REVIEW.parent.mkdir(parents=True, exist_ok=True)
    with REVIEW.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["source", "name", "what was done", "detail"])
        w.writerows(review)

    # Counts
    print(f"source rows: sheet {len(sheet_rows)} {dict(collections.Counter(x['src'] for x in sheet_rows))}")
    print(f"             notion live {len(live)} + archive {len(arch)}; skipped {len(skipped)}")
    print(f"total rows {len(rows)} -> people {len(cands)} (+{len(extra)} co-founders)")
    m = collections.Counter(c["match"] or "new" for c in cands)
    nobody = sum(1 for c in cands if c["nobody"])
    print(f"People: matched {sum(v for k, v in m.items() if k != 'new')} {dict(m)}; "
          f"new {m['new'] - nobody} + {len(extra)} co-founders; rows naming no founder {nobody}")
    print(f"People fills on existing rows: email {sum(c['fill_email'] for c in cands)}, LinkedIn {sum(c['fill_li'] for c in cands)}")
    print(f"Organisations: existing {sum(1 for c in cands if c['org_id'])}, new {len(new_orgs)}, none {sum(1 for c in cands if not c['company'])}")
    print(f"Engagements: existing {sum(1 for c in cands if c['eng_id'])}, new {sum(1 for c in cands if not c['eng_id'])}")
    print("status (new)", dict(collections.Counter(c["status"] for c in cands if not c["eng_id"])))
    print("host verdict", dict(collections.Counter(c["host"] for c in cands if not c["eng_id"])))
    print("tier", dict(collections.Counter(c["tier"] for c in cands if not c["eng_id"])))
    print("bucket", dict(collections.Counter(c["bucket"] for c in cands if not c["eng_id"])))
    print("waiting", dict(collections.Counter(",".join(c["waiting"]) for c in cands if not c["eng_id"])))
    print(f"review rows {len(review)}")


def cell(s):
    return str(s or "").replace("|", "/").replace("\n", " ").strip()


def write_prompt(cands, new_orgs, extra, existing):
    new_people = [c for c in cands if not c["person_id"] and not c["nobody"]]
    new_engs = [c for c in cands if not c["eng_id"]]
    upd_people = [c for c in cands if c["person_id"] and (c["fill_email"] or c["fill_li"])]
    upd_engs = [c for c in cands if c["eng_id"] and existing.get(c["eng_id"], {}).get("status") in ("Booked", "Contacted")]

    def person_ref(c):
        if c["nobody"]:
            return ""
        return f"https://app.notion.com/p/{c['person_id']}" if c["person_id"] else f"People row with Source ID \"{c['key']}\""

    def org_ref(c):
        if c["org_id"]:
            return f"https://app.notion.com/p/{c['org_id']}"
        if c["org_new"]:
            return f"Organisations row with Source ID \"{c['org_new']}\""
        return ""

    L = []
    L.append("# Built to Share pipeline: load into People, Organisations and Engagements")
    L.append("")
    L.append("You are working in Akshay's ThePodium Business OS. This file moves every Built to Share (B2S) guest list "
             "(Chaandini's Google Sheet and the old Built-2-Share Notion tables) into the three databases below. "
             "The research and matching are done: copy the tables exactly, do not re-match or guess.")
    L.append("")
    L.append("Databases:")
    L.append(f"- People: {PEOPLE_DB}")
    L.append(f"- Organisations: {ORGS_DB}")
    L.append(f"- Engagements: {ENGS_DB}")
    L.append(f"- Show for every Engagement: Built to Share, https://app.notion.com/p/{SHOW}")
    L.append("")
    L.append("## Rules")
    L.append("")
    L.append("1. Work through the parts in order: 1, 2, 3, 4, 5. A row in a later part links to rows made in an earlier one.")
    L.append("2. Before creating any row, check whether a row with the same Source ID already exists in that database. If it does, skip it. "
             "This makes the file safe to run again: if you stop part-way, carry on from where you stopped.")
    L.append("3. A cell that says `People row with Source ID \"X\"` or `Organisations row with Source ID \"X\"` means: find the row "
             "whose Source ID is exactly X and link it. A cell with a Notion URL means link that page.")
    L.append("4. Leave a property empty when its cell is empty. Never overwrite a value that is already set, except where a part says so.")
    L.append("5. Multi-select cells list values separated by commas. Create the option if it does not exist yet.")
    L.append(f"6. Set on every row you create: Source = Migration, Migration batch = {BATCH}, Captured = 2026-10-09.")
    L.append("7. Do not send email, do not edit any other row, do not delete anything.")
    L.append("8. At the end, report the counts in part 6.")
    L.append("")
    L.append(f"## Part 1: new Organisations ({len(new_orgs)} rows)")
    L.append("")
    L.append("Create one Organisations row per line. Type = Startup.")
    L.append("")
    L.append("| Source ID | Name | Domain |")
    L.append("|---|---|---|")
    for o in new_orgs.values():
        L.append(f"| {o['key']} | {cell(o['name'])} | {cell(o['domain'])} |")
    L.append("")
    np = len(new_people) + len(extra)
    L.append(f"## Part 2: new People ({np} rows)")
    L.append("")
    L.append("Create one People row per line. Tags = B2S guest pipeline (add the option). Organisation links as given.")
    L.append("")
    L.append("Where Duplicate of has a page, also set Match status = Possible duplicate and Duplicate check = "
             "\"Same name, different company (B2S migration)\". These are different people who share a name.")
    L.append("")
    L.append("| Source ID | Name | Email | LinkedIn | Role | Location | Organisation | Duplicate of |")
    L.append("|---|---|---|---|---|---|---|---|")
    for c in new_people:
        dup = f"https://app.notion.com/p/{c['dup_of']}" if c.get("dup_of") else ""
        L.append(f"| {cell(c['key'])} | {cell(c['name'])} | {cell(c['email'])} | {cell(li_url(c['li']))} | {cell(c['title'])} | "
                 f"{cell(c['country'])} | {cell(org_ref(c))} | {dup} |")
    for e in extra:
        ref = f"https://app.notion.com/p/{e['org_id']}" if e["org_id"] else (f"Organisations row with Source ID \"{e['org_new']}\"" if e["org_new"] else "")
        L.append(f"| {cell(e['source_id'])} | {cell(e['name'])} |  |  | Co-founder |  | {cell(ref)} |  |")
    L.append("")
    L.append(f"## Part 3: fill empty fields on existing People ({len(upd_people)} rows)")
    L.append("")
    L.append("Open each page and fill only the fields that are empty there. Do not change anything else.")
    L.append("")
    L.append("| Page | Name | Email | LinkedIn |")
    L.append("|---|---|---|---|")
    for c in upd_people:
        L.append(f"| https://app.notion.com/p/{c['person_id']} | {cell(c['name'])} | {cell(c['email'] if c['fill_email'] else '')} | "
                 f"{cell(li_url(c['li']) if c['fill_li'] else '')} |")
    L.append("")
    L.append(f"## Part 4: new Engagements ({len(new_engs)} rows)")
    L.append("")
    L.append("Create one Engagements row per line. For every row: Show = Built to Share, Direction = Outbound. "
             "Name = the Name cell. Guest, Organisation, Via and Stakeholders link as given. Tags and Waiting on are multi-selects. "
             "Claude note goes in the Claude note property.")
    L.append("")
    L.append("| Source ID | Name | Guest | Organisation | Status | Host verdict | Tier | Tags | Via | Waiting on | Next date | Next action | Chase step | Stakeholders | Claude note |")
    L.append("|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|")
    for c in new_engs:
        comp = re.sub(r"\s*\(.*?\)\s*", " ", c["company"] or "").strip()
        who = "Founder" if c["nobody"] else c["name"]
        name = f"{who} @ {comp}, Built to Share" if comp else f"{who}, Built to Share"
        via = f"https://app.notion.com/p/{SATISH}" if c["connect"] == "Satish" else ""
        st = ", ".join(f"https://app.notion.com/p/{v}" if k == "id" else f"People row with Source ID \"{v}\"" for k, v in c["stakeholders"])
        L.append(f"| {cell(c['key'])} | {cell(name)} | {cell(person_ref(c))} | {cell(org_ref(c))} | {c['status']} | {c['host']} | {c['tier']} | "
                 f"{', '.join(c['eng_tags'])} | {via} | {', '.join(c['waiting'])} | {cell(c['next_date'])} | {cell(c['next_action'])} | "
                 f"{c['chase'] or ''} | {cell(st)} | {cell(c['note'])} |")
    L.append("")
    L.append(f"## Part 5: existing Built to Share Engagements still open ({len(upd_engs)} rows)")
    L.append("")
    L.append("Fill only the properties that are empty on these pages. Add the tag to any tags already there. Do not change Status.")
    L.append("")
    L.append("| Page | Name | Host verdict | Tier | Add tag |")
    L.append("|---|---|---|---|---|")
    for c in upd_engs:
        host = "Yes" if existing[c["eng_id"]]["status"] == "Booked" else c["host"]   # booked means Satish said yes
        L.append(f"| https://app.notion.com/p/{c['eng_id']} | {cell(existing[c['eng_id']]['name'])} | {host} | {c['tier']} | {c['bucket'] or ''} |")
    L.append("")
    L.append("## Part 6: views and check")
    L.append("")
    L.append("Add these views to the Engagements database. All filter on Show contains Built to Share.")
    L.append("")
    L.append("| View | Layout | Filter | Sort |")
    L.append("|---|---|---|---|")
    L.append("| B2S: approved, not yet mailed | Table | Host verdict is Yes and Status is New | Tier ascending |")
    L.append("| B2S: outreach | Table | Status is not any of Unresponsive, They declined, We declined, Parked, Form filled | Next date ascending |")
    L.append("| B2S: for Satish | Table | Host verdict is Not asked, Status is New, Tier is Tier 1 or Tier 1.5 | Tier ascending |")
    L.append("| B2S: re-engage | Table | Status is Unresponsive and Next date is on or before today | Next date ascending |")
    L.append("")
    L.append(f"Then count rows with Migration batch = {BATCH} and report:")
    L.append("")
    L.append("| Database | Expected |")
    L.append("|---|---|")
    L.append(f"| Organisations | {len(new_orgs)} |")
    L.append(f"| People | {np} |")
    L.append(f"| Engagements | {len(new_engs)} |")
    L.append("")
    st = collections.Counter(c["status"] for c in new_engs)
    L.append("Engagements by Status: " + ", ".join(f"{k} {v}" for k, v in sorted(st.items(), key=lambda kv: -kv[1])) + ".")
    L.append("")
    L.append("Report any line you could not complete, with its Source ID and the reason.")
    (OUT / "notion_ai_b2s_pipeline.md").write_text("\n".join(L) + "\n")


if __name__ == "__main__":
    main()
