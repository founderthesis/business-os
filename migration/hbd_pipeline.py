"""Human by Design guest pipeline: Chaandini's "HBD new" sheet and the old HBD Notion table into People,
Organisations and Engagements, as one Notion AI prompt.

Usage:
  python migration/hbd_pipeline.py

Reads (git-ignored, data/raw/hbd/):
  hbd_new_2026-10-10.xlsx                   export of Google Sheet "HBD new" (Chaandini), 5 tabs
  notion_hbd_outreach_2026-10-10.json       the old Notion table "HBD Guest Outreach" (113 rows, 23 Sep import)
  notion_people_2026-10-10.json             People: id|name|tags|email|linkedin|other emails|org id
  notion_orgs_2026-10-10.json               Organisations: id|name|domain|website
  notion_engagements_2026-10-10.json        Engagements with a Guest: id|guest json|show json|status
Writes:
  data/out/hbd_pipeline/candidates.json
  data/out/hbd_pipeline/notion_ai_hbd_pipeline.md   the one prompt Akshay uploads to Notion AI (decision 91)
  data/review/hbd_pipeline_review.csv

How HBD differs from Built to Share (decision 105): Arjun decides after a guest has said yes, not before outreach.
"Approved" in the sheet means the guest agreed; "Arjun rejected" or "Decline" written in a follow-up cell is Arjun's
no. The "Rockstar prof" tab is Arjun's own list, so its rows are Host verdict Yes and already mailed (8 Oct meeting).
Existing HBD Engagements are never touched. Re-run safely; never edits the raw files.
"""
import collections
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from b2s_pipeline import (ACADEMIC, ENGS_DB, FREE_MAIL, ORGS_DB, PEOPLE_DB, cell, date_of, email_info,  # noqa: E402
                          fill_of, org_name_key, split_name, split_people, titlecase, val)
from paas import EMAIL, fold, name_key  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data/raw/hbd"
OUT = ROOT / "data/out/hbd_pipeline"
REVIEW = ROOT / "data/review/hbd_pipeline_review.csv"
TODAY = dt.date(2026, 10, 10)
BATCH = "hbd-pipeline-2026-10-10"
SHEET = "HBD new"
SHOW = "3ed19a301a7f815585e7e6c8accaab23"   # Human by Design
RED = set()                                  # red rows all say "Arjun rejected" in text; the fill alone is not used
# Same name as two People rows: the one that is this person (research)
PERSON_FIX = {"amit jain|luma": "416d4e37f730402c8b31411e2ff933e7"}   # the Luma AI row from the B2S load

RANK = {"New": 0, "Contacted": 1, "Unresponsive": 2, "LT chase": 3, "Interested": 4, "Decision pending": 5,
        "They declined": 6, "We declined": 6}
STATUS = {"unresponsive": "Unresponsive", "maybe later": "LT chase", "declined": "They declined", "delined": "They declined",
          "chasing": "Contacted", "approved": "Interested", "replied/chase": "Interested"}
REJECT = re.compile(r"arjun rejected|^decline$|arjun declined", re.I)
LATER_2027 = re.compile(r"reach ?out in 2027|early 2027", re.I)
BUCKET = {"Sheet1": "AI researcher", "V2": "AI researcher", "V3": "AI researcher", "Notion": "AI researcher",
          "Rockstar prof": "Rockstar prof", "Fortune 500": "Fortune 500 AI leader"}


def clean_email(text):
    found = [e.lower().strip(".") for e in EMAIL.findall(str(text or ""))]
    return found[0] if found else None


def first_org(text):
    """'MPI for Informatics (Director); Saarland University (Professor)' -> ('MPI for Informatics', 'Director')."""
    t = re.sub(r"'s (CEO|founder|co-?founder)\b", "", str(text or "").split(";")[0], flags=re.I).strip()
    m = re.match(r"^(.*?)\s*\((.*?)\)\s*$", t)
    return (m.group(1).strip(), m.group(2).strip()) if m else (t or None, None)


def rec(tab, sid, name):
    return {"tab": tab, "sid": sid, "name": name, "company": None, "email": None, "other_email": None, "role": None,
            "location": None, "raw_status": None, "texts": [], "sent": [], "next": None, "red": False, "notes": []}


def follow(r, lo, hi, x):
    for c in r[lo:hi]:
        d = date_of(c)
        if d and d.year == 2027 and d.month >= 8:
            d = d.replace(year=2026)      # typo: a 2026 follow-up typed as 2027 (V2, Follow Up 4)
        if d:
            x["sent"].append(d)
        elif val(c) and val(c).upper() not in ("MAIL", "LI", "BOTH"):
            x["texts"].append(val(c))


def read_sheet():
    wb = openpyxl.load_workbook(RAW / "hbd_new_2026-10-10.xlsx")
    out = []
    ws = wb["Sheet1"]
    for i, r in enumerate(ws.iter_rows(min_row=3), start=3):
        if not val(r[0]):
            continue
        x = rec("Sheet1", f"{SHEET}/Sheet1/{i}", val(r[0]))
        x["company"] = re.sub(r"\s*\(.*?\)\s*$", "", val(r[1])) or None
        x["location"], x["email"], x["raw_status"] = val(r[2]) or None, clean_email(val(r[3])), val(r[4]) or None
        x["red"] = fill_of(r[0]) in RED
        if fill_of(r[0]) in ("FFFFFF00", "FFF4CCCC"):
            x["notes"].append("Highlighted " + ("yellow" if fill_of(r[0]) == "FFFFFF00" else "light red") + " in the sheet, meaning not recorded")
        follow(r, 6, 10, x)
        out.append(x)
    ws = wb["V2"]
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        if not val(r[0]):
            continue
        x = rec("V2", f"{SHEET}/V2/{i}", val(r[0]))
        x["email"] = clean_email(val(r[1]))
        x["company"], x["role"] = first_org(val(r[2]))
        x["raw_status"], x["red"] = val(r[3]) or None, fill_of(r[0]) in RED
        follow(r, 5, 9, x)
        out.append(x)
    ws = wb["V3"]
    for i, r in enumerate(ws.iter_rows(min_row=3), start=3):
        if not val(r[0]):
            continue
        n, c = split_name(val(r[0]))
        x = rec("V3", f"{SHEET}/V3/{i}", n)
        x["company"] = c or val(r[1]) or None
        x["role"], x["email"], x["raw_status"] = val(r[2]) or None, clean_email(val(r[3])), val(r[4]) or None
        x["red"] = fill_of(r[0]) in RED
        if c and val(r[1]):
            x["notes"].append(f"Also at {val(r[1])}")
        follow(r, 6, 10, x)
        out.append(x)
    ws = wb["Rockstar prof"]
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        if not val(r[0]):
            continue
        x = rec("Rockstar prof", f"{SHEET}/Rockstar prof/{i}", val(r[0]))
        x["company"] = val(r[1]) or val(r[2]) or None
        if val(r[1]) and val(r[2]):
            x["notes"].append(f"Also at {val(r[2])}")
        x["location"], x["email"], x["raw_status"] = val(r[3]) or None, clean_email(val(r[4])), val(r[5]) or None
        out.append(x)
    ws = wb["Fortune 500"]
    for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
        if not val(r[0]):
            continue
        x = rec("Fortune 500", f"{SHEET}/Fortune 500/{i}", val(r[0]))
        x["company"], x["role"] = val(r[1]) or None, val(r[2]) or None
        x["email"], x["other_email"] = clean_email(val(r[3])), clean_email(val(r[4]))
        x["raw_status"] = val(r[5]) or None
        out.append(x)
    return out


def read_notion():
    out = []
    for r in json.load(open(RAW / "notion_hbd_outreach_2026-10-10.json")):
        x = rec("Notion", f"Notion HBD Guest Outreach/{r['id']}", r["name"] or "")
        x["company"], x["role"] = first_org(r.get("company"))
        x["email"], x["raw_status"] = clean_email(r.get("email")), r.get("status")
        if r.get("note"):
            x["texts"].append(r["note"])
        if r.get("next"):
            x["next"] = dt.date.fromisoformat(r["next"][:10])
        out.append(x)
    return out


def status_of(x):
    raw = (x["raw_status"] or "").strip().lower()
    s = STATUS.get(raw)
    if raw and s is None:
        return None
    if s is None:
        if x["tab"] == "Rockstar prof":
            s = "Contacted"          # Arjun's list was mailed before 8 Oct; the sheet keeps no dates for it
        else:
            s = "Contacted" if x["sent"] or x["next"] else "New"
    return s


def main():
    rows = read_sheet() + read_notion()
    review = []
    for x in rows:
        x["status"] = status_of(x)
        if x["status"] is None:
            review.append([x["sid"], x["name"], f"unmapped status {x['raw_status']!r}, read as New", ""])
            x["status"] = "New"
        txt = " | ".join(x["texts"])
        x["rejected"] = x["red"] or bool(REJECT.search(txt))
        x["later_2027"] = bool(LATER_2027.search(txt))
        x["no_call_yet"] = "no call yet" in txt.lower()
        x["people"] = [titlecase(p) for p in split_people(x["name"])] or [x["name"]]

    groups, by = [], {}
    for x in rows:
        keys = (["em:" + x["email"]] if x["email"] else []) + ["nm:" + name_key(x["people"][0])]
        g = next((by[k] for k in keys if k in by), None)
        if g is None:
            g = []
            groups.append(g)
        g.append(x)
        for k in keys:
            by.setdefault(k, g)

    P_name, P_email, P_org = collections.defaultdict(list), {}, {}
    for line in json.load(open(RAW / "notion_people_2026-10-10.json")):
        pid, nm, _tags, em, _li, other, org = (line.split("|") + [""] * 7)[:7]
        P_name[name_key(nm)].append(pid)
        for e in [em] + EMAIL.findall(other):
            if e:
                P_email.setdefault(e.lower(), pid)
        P_org[pid] = org
    O_name, O_dom, O_label = collections.defaultdict(list), {}, {}
    for line in json.load(open(RAW / "notion_orgs_2026-10-10.json")):
        oid, nm, dom, web = (line.split("|") + [""] * 4)[:4]
        O_name[org_name_key(nm)].append(oid)
        O_label[oid] = nm
        for d in (dom, re.sub(r"^https?://(www\.)?", "", web).split("/")[0]):
            if d:
                O_dom.setdefault(d.lower(), oid)
    hbd_eng = {}
    for line in json.load(open(RAW / "notion_engagements_2026-10-10.json")):
        eid, guest, show, status = (line.split("|") + [""] * 4)[:4]
        if SHOW in show:
            for gid in re.findall(r"[0-9a-f]{32}", guest):
                hbd_eng[gid] = (eid, status)

    tab_order = ["Rockstar prof", "V3", "V2", "Sheet1", "Fortune 500", "Notion"]
    cands, new_orgs = [], {}
    for xs in groups:
        xs.sort(key=lambda x: tab_order.index(x["tab"]))
        c = {"source_ids": [x["sid"] for x in xs], "tabs": list(dict.fromkeys(x["tab"] for x in xs))}
        c["name"] = max((x["people"][0] for x in xs), key=len)
        for f in ("company", "email", "other_email", "role", "location"):
            c[f] = next((x[f] for x in xs if x[f]), None)
        c["sent"] = sorted({d for x in xs for d in x["sent"]})
        c["next_src"] = max((x["next"] for x in xs if x["next"]), default=None)
        c["texts"] = list(dict.fromkeys(t for x in xs for t in x["texts"]))
        c["notes"] = list(dict.fromkeys(n for x in xs for n in x["notes"]))
        c["rejected"] = any(x["rejected"] for x in xs)
        c["arjun_list"] = "Rockstar prof" in c["tabs"]
        c["bucket"] = BUCKET[c["tabs"][0]]
        best = max((x["status"] for x in xs), key=lambda s: RANK[s])
        done = [d for d in c["sent"] if d <= TODAY]
        if best == "Contacted" and len(done) >= 4 and not [d for d in c["sent"] if d > TODAY] and done[-1] <= TODAY - dt.timedelta(days=7):
            best = "Unresponsive"
            c["notes"].insert(0, "Sequence finished with no reply.")
        if any(x["later_2027"] for x in xs) and RANK[best] < RANK["LT chase"]:
            best = "LT chase"
        if c["rejected"]:
            best = "We declined"
        elif any(x["no_call_yet"] for x in xs) and best == "Interested":
            best = "Decision pending"
        c["status"] = best
        c["raw_statuses"] = sorted({x["raw_status"] for x in xs if x["raw_status"]})
        cands.append(c)

    for c in cands:
        pid = PERSON_FIX.get(f"{name_key(c['name'])}|{fold(c['company'] or '')}")
        how = "research" if pid else None
        if pid:
            pass
        elif c["email"] and c["email"] in P_email:
            pid, how = P_email[c["email"]], "email"
        elif c["other_email"] and c["other_email"] in P_email:
            pid, how = P_email[c["other_email"]], "email"
        else:
            hits = P_name.get(name_key(c["name"]), []) if len(c["name"].split()) > 1 else []
            if len(hits) == 1:
                pid, how = hits[0], "name"
            elif len(hits) > 1:
                review.append([c["source_ids"][0], c["name"], "same name on several People rows; new row flagged Possible duplicate",
                               ", ".join(hits)])
                how = None
            else:
                how = None
        c["person_id"], c["match"] = pid, how
        c["existing"] = hbd_eng.get(pid) if pid else None
        c["dup_of"] = P_name[name_key(c["name"])][0] if not pid and len(P_name.get(name_key(c["name"]), [])) > 1 else None

        ok = org_name_key(c["company"]) if c["company"] else ""
        dom = c["email"].split("@")[1] if c["email"] else None
        if dom and (dom in FREE_MAIL or ACADEMIC.search(dom)):
            dom = None
        oid = O_name[ok][0] if ok and O_name.get(ok) else (O_dom.get(dom) if dom else None)
        c["org_id"], c["org_new"] = oid, None
        if not oid and c["company"]:
            if ok not in new_orgs:
                new_orgs[ok] = {"key": f"HBD org {len(new_orgs) + 1:03d}", "name": c["company"], "domain": dom,
                                "type": "Enterprise" if c["bucket"] == "Fortune 500 AI leader" else ""}
            c["org_new"] = new_orgs[ok]["key"]

        s = c["status"]
        host = "No" if c["rejected"] else ("Yes" if c["arjun_list"] else "Not asked")
        tier = "Tier 1.5" if c["bucket"] == "Fortune 500 AI leader" else "Tier 1"
        tags = [c["bucket"]] + ([] if c["email"] else ["No email"])
        last = c["sent"][-1] if c["sent"] else None
        nxt, action, waiting = None, None, []
        if s == "Contacted":
            future = [d for d in c["sent"] if d > TODAY]
            nxt = min(future) if future else max(TODAY, (last or c["next_src"] or TODAY) + dt.timedelta(days=7))
            action, waiting = f"Follow-up {len([d for d in c['sent'] if d <= TODAY]) + 1}", ["Guest or PR"]
        elif s == "Interested":
            nxt, action, waiting = TODAY, "Run the HBD eval, then ask Arjun", ["Chaandini"]
        elif s == "Decision pending":
            nxt, action, waiting = TODAY + dt.timedelta(days=3), "Arjun to decide", ["Arjun"]
        elif s == "LT chase":
            nxt = dt.date(2027, 1, 5) if any("2027" in t for t in c["texts"]) else (last or TODAY) + dt.timedelta(days=60)
            action, waiting = "Try again: they said later", ["Chaandini"]
        elif s == "Unresponsive":
            nxt, action = (last or TODAY) + dt.timedelta(days=60), "Re-engage: 2 months since the last mail"
        elif s == "New":
            action = "Ask Arjun" if tier == "Tier 1.5" else "First mail"
        note = [f"HBD migration 10 Oct, from {', '.join(c['tabs'])}."]
        if c["raw_statuses"]:
            note.append(f"Was: {', '.join(c['raw_statuses'])}.")
        if c["texts"]:
            note.append("Sheet notes: " + "; ".join(t[:80] for t in c["texts"][:3]) + ".")
        if c["sent"]:
            note.append(f"Mails sent: {len([d for d in c['sent'] if d <= TODAY])}, last {last:%d %b}.")
        if c["arjun_list"]:
            note.append("On Arjun's own list (Rockstar prof).")
        if c["rejected"]:
            note.append("Arjun said no.")
        note += c["notes"][:2]
        c.update(host=host, tier=tier, tags=tags, next_date=nxt.isoformat() if nxt else None, next_action=action,
                 waiting=waiting, chase=len([d for d in c["sent"] if d <= TODAY]) or None, note=" ".join(note)[:500])

    for c in cands:
        c["key"] = c["source_ids"][0]
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "candidates.json").write_text(json.dumps({"cands": cands, "orgs": new_orgs}, indent=1, default=str))
    for c in cands:
        if c["existing"]:
            review.append([c["key"], c["name"], f"already an HBD Engagement ({c['existing'][1]}); left as it is", c["existing"][0]])
        if c["match"] == "name" and c["company"] and P_org.get(c["person_id"]) and c["org_id"] and c["org_id"] != P_org[c["person_id"]]:
            review.append([c["key"], c["name"], "matched People by name; Organisation differs there", c["company"]])
    REVIEW.parent.mkdir(parents=True, exist_ok=True)
    with REVIEW.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["source", "name", "what was done", "detail"])
        w.writerows(review)
    write_prompt(cands, new_orgs)

    new_engs = [c for c in cands if not c["existing"]]
    print(f"source rows {len(rows)} {dict(collections.Counter(x['tab'] for x in rows))}")
    print(f"people {len(cands)}; matched {sum(1 for c in cands if c['person_id'])} {dict(collections.Counter(c['match'] for c in cands))}")
    print(f"orgs: existing {sum(1 for c in cands if c['org_id'])}, new {len(new_orgs)}")
    print(f"engagements: existing HBD {sum(1 for c in cands if c['existing'])}, new {len(new_engs)}")
    print("status", dict(collections.Counter(c["status"] for c in new_engs)))
    print("host", dict(collections.Counter(c["host"] for c in new_engs)), "tier", dict(collections.Counter(c["tier"] for c in new_engs)))
    print("bucket", dict(collections.Counter(c["bucket"] for c in new_engs)))
    print("waiting", dict(collections.Counter(",".join(c["waiting"]) for c in new_engs)))
    print(f"review rows {len(review)}")


def write_prompt(cands, new_orgs):
    new_people = [c for c in cands if not c["person_id"]]
    new_engs = [c for c in cands if not c["existing"]]

    def org_ref(c):
        if c["org_id"]:
            return f"https://app.notion.com/p/{c['org_id']}"
        return f"Organisations row with Source ID \"{c['org_new']}\"" if c["org_new"] else ""

    L = ["# Human by Design pipeline: load into People, Organisations and Engagements", "",
         "You are working in Akshay's ThePodium Business OS. This file moves Chaandini's Human by Design (HBD) guest lists "
         "(her \"HBD new\" Google Sheet and the old \"HBD Guest Outreach\" Notion table) into the three databases below. "
         "The research and matching are done: copy the tables exactly, do not re-match or guess.", "",
         "Databases:", f"- People: {PEOPLE_DB}", f"- Organisations: {ORGS_DB}", f"- Engagements: {ENGS_DB}",
         f"- Show for every Engagement: Human by Design, https://app.notion.com/p/{SHOW}", "",
         "## Rules", "",
         "1. Work through the parts in order. A row in a later part links to rows made in an earlier one.",
         "2. Before creating any row, check whether a row with the same Source ID already exists in that database. If it does, skip it. "
         "If you stop part-way, carry on from where you stopped.",
         "3. A cell that says `People row with Source ID \"X\"` or `Organisations row with Source ID \"X\"` means: link the row whose "
         "Source ID is exactly X. A cell with a Notion URL means link that page.",
         "4. Leave a property empty when its cell is empty. Multi-select cells list values separated by commas; create an option if it is missing.",
         f"5. Set on every row you create: Source = Migration, Migration batch = {BATCH}, Captured = 2026-10-10.",
         "6. Do not edit any existing row, do not send email, do not delete anything.", "",
         f"## Part 1: new Organisations ({len(new_orgs)} rows)", "",
         "| Source ID | Name | Domain | Type |", "|---|---|---|---|"]
    for o in new_orgs.values():
        L.append(f"| {o['key']} | {cell(o['name'])} | {cell(o['domain'])} | {o['type']} |")
    L += ["", f"## Part 2: new People ({len(new_people)} rows)", "",
          "Tags = HBD guest pipeline (add the option). Where Duplicate of has a page, also set Match status = Possible duplicate.", "",
          "| Source ID | Name | Email | Other emails | Role | Location | Organisation | Duplicate of |", "|---|---|---|---|---|---|---|---|"]
    for c in new_people:
        dup = f"https://app.notion.com/p/{c['dup_of']}" if c["dup_of"] else ""
        L.append(f"| {cell(c['key'])} | {cell(c['name'])} | {cell(c['email'])} | {cell(c['other_email'])} | {cell(c['role'])} | "
                 f"{cell(c['location'])} | {cell(org_ref(c))} | {dup} |")
    L += ["", f"## Part 3: new Engagements ({len(new_engs)} rows)", "",
          "For every row: Show = Human by Design, Direction = Outbound. Guest and Organisation link as given.", "",
          "| Source ID | Name | Guest | Organisation | Status | Host verdict | Tier | Tags | Waiting on | Next date | Next action | Chase step | Claude note |",
          "|---|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for c in new_engs:
        guest = f"https://app.notion.com/p/{c['person_id']}" if c["person_id"] else f"People row with Source ID \"{c['key']}\""
        name = f"{c['name']} @ {c['company']}, Human by Design" if c["company"] else f"{c['name']}, Human by Design"
        L.append(f"| {cell(c['key'])} | {cell(name)} | {guest} | {cell(org_ref(c))} | {c['status']} | {c['host']} | {c['tier']} | "
                 f"{', '.join(c['tags'])} | {', '.join(c['waiting'])} | {cell(c['next_date'])} | {cell(c['next_action'])} | "
                 f"{c['chase'] or ''} | {cell(c['note'])} |")
    st = collections.Counter(c["status"] for c in new_engs)
    L += ["", "## Part 4: check", "", f"Count rows with Migration batch = {BATCH} and report:", "",
          "| Database | Expected |", "|---|---|", f"| Organisations | {len(new_orgs)} |", f"| People | {len(new_people)} |",
          f"| Engagements | {len(new_engs)} |", "",
          "Engagements by Status: " + ", ".join(f"{k} {v}" for k, v in st.most_common()) + ".", "",
          "Report any line you could not complete, with its Source ID and the reason."]
    (OUT / "notion_ai_hbd_pipeline.md").write_text("\n".join(L) + "\n")


if __name__ == "__main__":
    main()
